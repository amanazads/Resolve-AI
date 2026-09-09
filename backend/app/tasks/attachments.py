"""
Generic multi-format attachment & document processing layer for Resolve AI.

Parses files (CSV, XLSX, PDF, DOCX, TXT, JSON, MD) into structured representations
and concise text contexts for agent multi-file reasoning.
"""

import csv
import hashlib
import io
import json
import logging
import mimetypes
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.config import settings
from app.tasks.models import TaskArtifact

logger = logging.getLogger(__name__)

DANGEROUS_EXTENSIONS = {
    ".exe", ".bat", ".cmd", ".sh", ".bash", ".bin", ".msi", ".dll", ".so", ".dmg",
    ".app", ".vbs", ".ps1", ".jar", ".scr", ".pif", ".cpl", ".iso"
}


class AttachmentProcessor:
    """
    Ingests, validates, and extracts text/data from diverse file formats with security controls.
    """

    @classmethod
    def process_file(
        cls,
        content: bytes,
        filename: str,
        metadata: Optional[Dict[str, Any]] = None,
        task_id: Optional[str] = None,
    ) -> TaskArtifact:
        """
        Parses raw bytes into a populated, secure TaskArtifact.
        """
        size_bytes = len(content)
        max_bytes = getattr(settings, "MAX_FILE_SIZE_BYTES", 25 * 1024 * 1024)
        if size_bytes > max_bytes:
            raise ValueError(
                f"File '{filename}' ({size_bytes} bytes) exceeds maximum permitted size of {max_bytes} bytes."
            )

        suffix = Path(filename).suffix.lower()
        if suffix in DANGEROUS_EXTENSIONS:
            raise ValueError(f"File '{filename}' has dangerous or executable extension '{suffix}' and is rejected.")

        sha256_hash = hashlib.sha256(content).hexdigest()
        guessed_mime, _ = mimetypes.guess_type(filename)
        mime_type = guessed_mime or "application/octet-stream"

        lower_name = filename.lower()
        meta = dict(metadata or {})
        extracted_text = ""
        parsed_data = None
        file_type = "unknown"

        try:
            if lower_name.endswith(".csv"):
                file_type = "csv"
                mime_type = mime_type or "text/csv"
                extracted_text, parsed_data = cls._parse_csv(content)
            elif lower_name.endswith((".xlsx", ".xlsm")):
                file_type = "xlsx"
                mime_type = mime_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                extracted_text, parsed_data = cls._parse_xlsx(content)
            elif lower_name.endswith(".pdf"):
                file_type = "pdf"
                mime_type = mime_type or "application/pdf"
                extracted_text = cls._parse_pdf(content)
            elif lower_name.endswith((".docx", ".doc")):
                file_type = "docx"
                mime_type = mime_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                extracted_text = cls._parse_docx(content)
            elif lower_name.endswith(".json"):
                file_type = "json"
                mime_type = mime_type or "application/json"
                extracted_text, parsed_data = cls._parse_json(content)
            else:
                file_type = "text"
                mime_type = mime_type or "text/plain"
                extracted_text = cls._parse_text(content)

        except Exception as e:
            logger.warning("Error processing attachment '%s': %s", filename, e)
            extracted_text = f"[Error reading file content: {e}]"

        # Build clean summary preview
        summary = (
            extracted_text[:2000] + ("..." if len(extracted_text) > 2000 else "")
            if extracted_text
            else ""
        )
        meta["preview"] = summary
        meta["sha256"] = sha256_hash

        return TaskArtifact(
            filename=filename,
            file_type=file_type,
            mime_type=mime_type,
            size_bytes=size_bytes,
            sha256_hash=sha256_hash,
            extracted_text=extracted_text,
            structured_data=parsed_data,
            parsed_data=parsed_data,
            task_id=task_id,
            metadata=meta,
        )

    @classmethod
    def _parse_csv(cls, content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        text = content.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        headers = reader.fieldnames or []

        summary_lines = [
            f"CSV file with {len(rows)} rows and columns: {', '.join(headers)}",
            "Sample rows (up to 10):",
        ]
        for idx, row in enumerate(rows[:10]):
            clean_row = {k: v for k, v in row.items() if v}
            summary_lines.append(f"  Row {idx + 1}: {json.dumps(clean_row)}")

        return "\n".join(summary_lines), rows

    @classmethod
    def _parse_xlsx(cls, content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return "Empty Excel workbook", []

        headers = [str(cell) if cell is not None else f"col_{idx}" for idx, cell in enumerate(rows[0])]
        parsed_rows = []
        for row in rows[1:]:
            row_dict = {}
            for col_idx, cell in enumerate(row):
                if col_idx < len(headers) and cell is not None:
                    row_dict[headers[col_idx]] = str(cell)
            if row_dict:
                parsed_rows.append(row_dict)

        summary_lines = [
            f"Excel file with {len(parsed_rows)} rows. Sheet '{sheet.title}'. Columns: {', '.join(headers)}",
            "Sample rows (up to 10):",
        ]
        for idx, row in enumerate(parsed_rows[:10]):
            summary_lines.append(f"  Row {idx + 1}: {json.dumps(row)}")

        return "\n".join(summary_lines), parsed_rows

    @classmethod
    def _parse_pdf(cls, content: bytes) -> str:
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                if txt.strip():
                    pages_text.append(f"--- Page {idx + 1} ---\n{txt.strip()}")
            return "\n\n".join(pages_text) if pages_text else "[PDF contains no readable text]"
        except ImportError:
            return "[pypdf not installed]"

    @classmethod
    def _parse_docx(cls, content: bytes) -> str:
        try:
            import docx

            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs) if paragraphs else "[DOCX contains no readable paragraphs]"
        except ImportError:
            return "[python-docx not installed]"

    @classmethod
    def _parse_json(cls, content: bytes) -> Tuple[str, Any]:
        text = content.decode("utf-8", errors="replace")
        data = json.loads(text)
        return f"JSON data:\n{json.dumps(data, indent=2)[:5000]}", data

    @classmethod
    def _parse_text(cls, content: bytes) -> str:
        return content.decode("utf-8", errors="replace")
